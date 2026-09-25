from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from pathlib import Path
import re
import json

BASE = Path(__file__).resolve().parent.parent
FRONTEND = BASE / "frontend"

app = Flask(__name__)
CORS(app)

SKILLS = {
    "Python": ["python"], "Java": ["java"], "C++": ["c++", "cpp"],
    "JavaScript": ["javascript", "js"], "TypeScript": ["typescript", "ts"],
    "React": ["react", "reactjs"], "Node.js": ["node.js", "nodejs"],
    "Flask": ["flask"], "Django": ["django"], "SQL": ["sql"],
    "MySQL": ["mysql"], "MongoDB": ["mongodb"], "Git": ["git"],
    "GitHub": ["github"], "HTML": ["html"], "CSS": ["css"],
    "Machine Learning": ["machine learning", "ml"], "Deep Learning": ["deep learning"],
    "NLP": ["nlp", "natural language processing"], "Data Science": ["data science"],
    "Pandas": ["pandas"], "NumPy": ["numpy"], "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"], "Power BI": ["power bi"], "Excel": ["excel"],
    "AWS": ["aws", "amazon web services"], "Docker": ["docker"],
    "REST API": ["rest api", "restful"], "Communication": ["communication"],
}

def extract_text(file):
    name = file.filename.lower()
    raw = file.read()
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            from io import BytesIO
            reader = PdfReader(BytesIO(raw))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:
            raise RuntimeError("PDF reading failed. Install dependencies with START.bat.")
    if name.endswith(".docx"):
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            raise RuntimeError("DOCX reading failed. Install dependencies with START.bat.")
    raise RuntimeError("Supported formats: PDF, DOCX, TXT")

def find_skills(text):
    low = text.lower()
    found = []
    for display, variants in SKILLS.items():
        if any(re.search(r"(?<!\w)" + re.escape(v) + r"(?!\w)", low) for v in variants):
            found.append(display)
    return found

def extract_email(text):
    m = re.search(r'[\w.+-]+@[\w-]+\.[\w.-]+', text)
    return m.group(0) if m else ""

def extract_phone(text):
    m = re.search(r'(?<!\d)(?:\+91[-\s]?)?[6-9]\d{9}(?!\d)', re.sub(r'[()]', '', text))
    return m.group(0) if m else ""

def score_resume(text, skills):
    words = len(re.findall(r"\b\w+\b", text))
    sections = ["education", "experience", "project", "skills", "summary", "contact"]
    section_hits = sum(1 for s in sections if re.search(r"\b"+s+r"\b", text, re.I))
    skill_score = min(45, len(skills) * 3)
    section_score = min(30, section_hits * 5)
    length_score = 25 if 300 <= words <= 900 else (15 if 180 <= words <= 1200 else 8)
    return min(100, skill_score + section_score + length_score)

def jd_match(resume_skills, jd_text):
    jd_skills = find_skills(jd_text)
    matched = sorted(set(resume_skills) & set(jd_skills))
    missing = sorted(set(jd_skills) - set(resume_skills))
    pct = round(len(matched) / len(jd_skills) * 100) if jd_skills else 0
    return jd_skills, matched, missing, pct

@app.get("/")
def home():
    return send_from_directory(FRONTEND, "index.html")

@app.get("/<path:path>")
def static_files(path):
    return send_from_directory(FRONTEND, path)

@app.post("/api/analyse")
def analyse():
    if "resume" not in request.files:
        return jsonify({"error": "Please upload a resume."}), 400
    f = request.files["resume"]
    try:
        text = extract_text(f)
        skills = find_skills(text)
        result = {
            "name": (re.search(r"(?im)^(?:name\s*[:\-]\s*)?([A-Z][A-Za-z .'-]{2,50})$", text.strip()) or [None, ""])[1].strip(),
            "email": extract_email(text),
            "phone": extract_phone(text),
            "skills": skills,
            "skill_count": len(skills),
            "ats_score": score_resume(text, skills),
            "word_count": len(re.findall(r"\b\w+\b", text)),
            "suggestions": []
        }
        if result["ats_score"] < 70:
            result["suggestions"].append("Add clear sections: Summary, Skills, Education, Experience and Projects.")
        if len(skills) < 8:
            result["suggestions"].append("Add relevant technical skills that you genuinely know.")
        if not re.search(r"\b(projects?|experience)\b", text, re.I):
            result["suggestions"].append("Add 2–3 strong project or experience entries with measurable outcomes.")
        if not result["email"]:
            result["suggestions"].append("Add a professional email address.")
        if not result["phone"]:
            result["suggestions"].append("Add a valid phone number.")
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.post("/api/match")
def match():
    if "resume" not in request.files or "job_description" not in request.form:
        return jsonify({"error": "Resume and job description are required."}), 400
    try:
        text = extract_text(request.files["resume"])
        skills = find_skills(text)
        jd_skills, matched, missing, pct = jd_match(skills, request.form["job_description"])
        return jsonify({"job_skills": jd_skills, "matched": matched, "missing": missing, "match_percentage": pct})
    except Exception as e:
        return jsonify({"error": str(e)}), 400
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from pathlib import Path
import re
import json

BASE = Path(__file__).resolve().parent.parent
FRONTEND = BASE / "frontend"

app = Flask(__name__)
CORS(app)

SKILLS = {
    "Python": ["python"], "Java": ["java"], "C++": ["c++", "cpp"],
    "JavaScript": ["javascript", "js"], "TypeScript": ["typescript", "ts"],
    "React": ["react", "reactjs"], "Node.js": ["node.js", "nodejs"],
    "Flask": ["flask"], "Django": ["django"], "SQL": ["sql"],
    "MySQL": ["mysql"], "MongoDB": ["mongodb"], "Git": ["git"],
    "GitHub": ["github"], "HTML": ["html"], "CSS": ["css"],
    "Machine Learning": ["machine learning", "ml"], "Deep Learning": ["deep learning"],
    "NLP": ["nlp", "natural language processing"], "Data Science": ["data science"],
    "Pandas": ["pandas"], "NumPy": ["numpy"], "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"], "Power BI": ["power bi"], "Excel": ["excel"],
    "AWS": ["aws", "amazon web services"], "Docker": ["docker"],
    "REST API": ["rest api", "restful"], "Communication": ["communication"],
}

def extract_text(file):
    name = file.filename.lower()
    raw = file.read()
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            from io import BytesIO
            reader = PdfReader(BytesIO(raw))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:
            raise RuntimeError("PDF reading failed. Install dependencies with START.bat.")
    if name.endswith(".docx"):
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            raise RuntimeError("DOCX reading failed. Install dependencies with START.bat.")
    raise RuntimeError("Supported formats: PDF, DOCX, TXT")

def find_skills(text):
    low = text.lower()
    found = []
    for display, variants in SKILLS.items():
        if any(re.search(r"(?<!\w)" + re.escape(v) + r"(?!\w)", low) for v in variants):
            found.append(display)
    return found

def extract_email(text):
    m = re.search(r'[\w.+-]+@[\w-]+\.[\w.-]+', text)
    return m.group(0) if m else ""

def extract_phone(text):
    m = re.search(r'(?<!\d)(?:\+91[-\s]?)?[6-9]\d{9}(?!\d)', re.sub(r'[()]', '', text))
    return m.group(0) if m else ""

def score_resume(text, skills):
    words = len(re.findall(r"\b\w+\b", text))
    sections = ["education", "experience", "project", "skills", "summary", "contact"]
    section_hits = sum(1 for s in sections if re.search(r"\b"+s+r"\b", text, re.I))
    skill_score = min(45, len(skills) * 3)
    section_score = min(30, section_hits * 5)
    length_score = 25 if 300 <= words <= 900 else (15 if 180 <= words <= 1200 else 8)
    return min(100, skill_score + section_score + length_score)

def jd_match(resume_skills, jd_text):
    jd_skills = find_skills(jd_text)
    matched = sorted(set(resume_skills) & set(jd_skills))
    missing = sorted(set(jd_skills) - set(resume_skills))
    pct = round(len(matched) / len(jd_skills) * 100) if jd_skills else 0
    return jd_skills, matched, missing, pct

@app.get("/")
def home():
    return send_from_directory(FRONTEND, "index.html")

@app.get("/<path:path>")
def static_files(path):
    return send_from_directory(FRONTEND, path)

@app.post("/api/analyse")
def analyse():
    if "resume" not in request.files:
        return jsonify({"error": "Please upload a resume."}), 400
    f = request.files["resume"]
    try:
        text = extract_text(f)
        skills = find_skills(text)
        result = {
            "name": (re.search(r"(?im)^(?:name\s*[:\-]\s*)?([A-Z][A-Za-z .'-]{2,50})$", text.strip()) or [None, ""])[1].strip(),
            "email": extract_email(text),
            "phone": extract_phone(text),
            "skills": skills,
            "skill_count": len(skills),
            "ats_score": score_resume(text, skills),
            "word_count": len(re.findall(r"\b\w+\b", text)),
            "suggestions": []
        }
        if result["ats_score"] < 70:
            result["suggestions"].append("Add clear sections: Summary, Skills, Education, Experience and Projects.")
        if len(skills) < 8:
            result["suggestions"].append("Add relevant technical skills that you genuinely know.")
        if not re.search(r"\b(projects?|experience)\b", text, re.I):
            result["suggestions"].append("Add 2–3 strong project or experience entries with measurable outcomes.")
        if not result["email"]:
            result["suggestions"].append("Add a professional email address.")
        if not result["phone"]:
            result["suggestions"].append("Add a valid phone number.")
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.post("/api/match")
def match():
    if "resume" not in request.files or "job_description" not in request.form:
        return jsonify({"error": "Resume and job description are required."}), 400
    try:
        text = extract_text(request.files["resume"])
        skills = find_skills(text)
        jd_skills, matched, missing, pct = jd_match(skills, request.form["job_description"])
        return jsonify({"job_skills": jd_skills, "matched": matched, "missing": missing, "match_percentage": pct})
    except Exception as e:
        return jsonify({"error": str(e)}), 400
        
if __name__ == "__main__":
    import os
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
